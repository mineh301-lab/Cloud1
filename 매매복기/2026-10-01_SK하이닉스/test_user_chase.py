#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""사용자님 1개월 주문내역(9/8~9/28)에서 하이닉스 매수 체결을 정정 이력별로 나눠 당일 종가 대비 성과를 본다."""
import logging
import sys
logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s", stream=sys.stdout)
log = logging.getLogger("chase")
import json
import datetime as dt
from collections import defaultdict
import openpyxl

MONTH_XLSX = sys.argv[1]
DAILY_JSON = sys.argv[2] if len(sys.argv) > 2 else "data/yd_000660.json"


def load_close():
    j = json.load(open(DAILY_JSON))["chart"]["result"][0]
    out = {}
    for i, t in enumerate(j["timestamp"]):
        d = (dt.datetime.utcfromtimestamp(t) + dt.timedelta(hours=9)).date()
        out[d] = j["indicators"]["quote"][0]["close"][i]
    return out


def main():
    ws = openpyxl.load_workbook(MONTH_XLSX, data_only=True).worksheets[0]
    recs, r = [], 3
    while r <= ws.max_row:
        a = [ws.cell(row=r, column=c).value for c in range(1, 12)]
        b = [ws.cell(row=r + 1, column=c).value for c in range(1, 12)]
        r += 2
        if a[0] is None and b[0] is None:
            continue
        recs.append(dict(no=a[1], orig=a[2], code=a[3], side=a[5] or "", price=a[7], date=b[0].date() if b[0] else None,
                         fillq=b[4] or 0, fillp=b[5], mod=b[6], ftime=b[9], ex=b[10]))
    hy = [x for x in recs if x["code"] == "000660"]
    by_no = {x["no"]: x for x in hy}
    kids = defaultdict(list)
    for x in hy:
        if x["orig"]:
            kids[x["orig"]].append(x["no"])

    def chain(no):
        out = [by_no[no]]
        for c in sorted(kids.get(no, [])):
            out += chain(c)
        return out

    close = load_close()
    rows = []
    for root in [x for x in hy if not x["orig"] and "매수" in x["side"]]:
        ch = chain(root["no"])
        fills = [c for c in ch if c["fillq"] > 0]
        if not fills:
            continue
        prices = [c["price"] for c in ch if c["price"]]
        ups = sum(1 for i in range(1, len(prices)) if prices[i] > prices[i - 1])
        for f in fills:
            cl = close.get(f["date"])
            if cl is None:
                log.warning(f"[CHASE] {f['date']} 종가 없음 — 제외")
                continue
            ret = cl / f["fillp"] - 1
            pre = f["ftime"] is not None and f["ftime"] < dt.time(9, 0)
            post = f["ftime"] is not None and f["ftime"] > dt.time(15, 30)
            rows.append(dict(date=f["date"], t=f["ftime"], px=f["fillp"], q=f["fillq"], ups=ups, ret=ret, pre=pre, post=post))
            log.info(f"[FILL] {f['date']} {f['ftime']} {f['fillq']}@{f['fillp']:,} | 위로정정 {ups}회 | 당일종가 {cl:,.0f} | {ret*100:+.2f}%"
                     f"{' (장전)' if pre else ''}{' (장후)' if post else ''}")
    mean = lambda xs: sum(xs) / len(xs) if xs else float("nan")
    for tag, sub in [("위로 정정 후 체결", [x for x in rows if x["ups"] > 0]),
                     ("정정 없이 체결", [x for x in rows if x["ups"] == 0]),
                     ("위로 정정 후 체결(정규장만)", [x for x in rows if x["ups"] > 0 and not x["pre"] and not x["post"]]),
                     ("정정 없이 체결(정규장만)", [x for x in rows if x["ups"] == 0 and not x["pre"] and not x["post"]])]:
        log.info(f"[CHASESUM] {tag} | 체결 {len(sub)}건 | 당일종가 대비 평균 {mean([x['ret'] for x in sub])*100:+.2f}% | "
                 f"종가가 체결가 위인 비율 {mean([1.0 if x['ret'] > 0 else 0.0 for x in sub]):.0%}")


if __name__ == "__main__":
    main()
